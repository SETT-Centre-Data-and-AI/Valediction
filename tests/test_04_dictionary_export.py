from contextlib import contextmanager
from pathlib import Path

import pytest
from openpyxl import load_workbook
from support import compare_dictionary_columns, compare_dictionary_tables

from valediction.demo.demo_dictionary import demo_dictionary
from valediction.dictionary.importing import import_dictionary

EXPORT_DIR = Path(__file__).parent
EXPORT_FILENAME = "TEST - Data Dictionary"


@contextmanager
def cleanup():
    path = EXPORT_DIR / f"{EXPORT_FILENAME}.xlsx"
    if path.exists():
        path.unlink()
    try:
        yield path  # caller uses this path
    finally:
        try:
            path.unlink()  # post-clean
        except FileNotFoundError:
            pass


def test_dictionary_export() -> None:
    with cleanup():
        # Export Data Dictionary
        dictionary = demo_dictionary()
        dictionary.export_dictionary(directory=EXPORT_DIR, filename=EXPORT_FILENAME)


def test_dictionary_export_format() -> None:
    with cleanup():
        # Export Demo Data Dictionary
        dictionary = demo_dictionary()
        dictionary["DIAGNOSES"]["PATIENT_HASH"].vocabulary = ["ICD-10", "SNOMED"]
        dictionary["DIAGNOSES"]["PATIENT_HASH"].foreign_key = [
            "DEMOGRAPHICS.PATIENT_HASH",
            "LAB_TESTS.PATIENT_HASH",
        ]
        dictionary["DIAGNOSES"]["PATIENT_HASH"].required = True
        dictionary.export_dictionary(directory=EXPORT_DIR, filename=EXPORT_FILENAME)

        workbook = load_workbook(EXPORT_DIR / f"{EXPORT_FILENAME}.xlsx")
        worksheet = next(
            sheet for sheet in workbook.worksheets if "columns" in sheet.title.lower()
        )
        header_row = next(
            row_index
            for row_index, row in enumerate(
                worksheet.iter_rows(values_only=True), start=1
            )
            if any(str(value).strip().lower() == "required" for value in row)
        )
        required_column = next(
            index
            for index in range(1, worksheet.max_column + 1)
            if str(worksheet.cell(header_row, index).value).strip().lower()
            == "required"
        )
        required_values = [
            worksheet.cell(row_index, required_column).value
            for row_index in range(header_row + 1, worksheet.max_row + 1)
        ]
        assert "Y" in required_values
        assert "N" not in required_values
        workbook.close()

        # Reimport Demo Data Dictionary
        exported = import_dictionary(EXPORT_DIR / f"{EXPORT_FILENAME}.xlsx")

        # Check Match
        compare_dictionary_tables(exported, dictionary)
        compare_dictionary_columns(exported, dictionary)


def test_dictionary_export_without_metadata() -> None:
    with cleanup():
        # Export Data Dictionary
        dictionary = demo_dictionary()
        dictionary.name = None
        dictionary.version = None
        dictionary.inclusion_criteria = None
        dictionary.exclusion_criteria = None
        dictionary.version_notes = None
        dictionary.export_dictionary(
            directory=EXPORT_DIR, filename=EXPORT_FILENAME, overwrite=True
        )

        # Reimport Demo Data Dictionary
        exported = import_dictionary(EXPORT_DIR / f"{EXPORT_FILENAME}.xlsx")

        # Check Match
        compare_dictionary_tables(exported, dictionary)
        compare_dictionary_columns(exported, dictionary)


if __name__ == "__main__":
    pytest.main([__file__])
