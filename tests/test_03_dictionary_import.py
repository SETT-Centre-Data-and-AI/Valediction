import pytest
from openpyxl import load_workbook
from support import compare_dictionary_columns, compare_dictionary_tables

from valediction.demo import DEMO_DICTIONARY
from valediction.demo.demo_dictionary import demo_dictionary
from valediction.dictionary.exporting import export_dictionary
from valediction.dictionary.importing import import_dictionary
from valediction.dictionary.model import Dictionary
from valediction.exceptions import DataDictionaryImportError


def check_dictionary() -> Dictionary:
    return demo_dictionary()


def imported_dictionary() -> Dictionary:
    return import_dictionary(DEMO_DICTIONARY)


def _dictionary_with_metadata_cells(
    tmp_path,
    vocabulary: str,
    foreign_key: str,
    required: str | None = None,
):
    dictionary_path = tmp_path / "metadata.xlsx"
    export_dictionary(demo_dictionary(), tmp_path, filename="metadata")

    workbook = load_workbook(dictionary_path)
    worksheet = next(
        sheet for sheet in workbook.worksheets if "columns" in sheet.title.lower()
    )
    header_row = next(
        row_index
        for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1)
        if {"table", "column", "vocabularies", "foreign key target"}
        <= {str(value).strip().lower() for value in row if value is not None}
    )
    headers = {
        str(worksheet.cell(header_row, column_index).value)
        .strip()
        .lower(): column_index
        for column_index in range(1, worksheet.max_column + 1)
        if worksheet.cell(header_row, column_index).value is not None
    }
    for row_index in range(header_row + 1, worksheet.max_row + 1):
        table = worksheet.cell(row_index, headers["table"]).value
        column = worksheet.cell(row_index, headers["column"]).value
        if str(table).lower() == "diagnoses" and str(column).lower() == "patient_hash":
            worksheet.cell(row_index, headers["vocabularies"]).value = vocabulary
            worksheet.cell(row_index, headers["foreign key target"]).value = foreign_key
            if required is not None:
                worksheet.cell(row_index, headers["required"]).value = required
            break
    else:
        raise AssertionError("Could not find diagnoses.patient_hash in workbook")

    workbook.save(dictionary_path)
    return dictionary_path


def test_demo_dictionary_import():
    imported_dictionary = import_dictionary(DEMO_DICTIONARY)
    assert imported_dictionary
    assert all(not column.required for table in imported_dictionary for column in table)


def test_tables_imported_correctly():
    test_dictionary = imported_dictionary()
    check_dictionary = demo_dictionary()

    compare_dictionary_tables(test_dictionary, check_dictionary)


def test_columns_imported_correctly():
    test_dictionary = imported_dictionary()
    check_dictionary = demo_dictionary()

    compare_dictionary_columns(test_dictionary, check_dictionary)


def test_import_parses_multiple_metadata_values(tmp_path):
    path = _dictionary_with_metadata_cells(
        tmp_path,
        vocabulary=" ICD-10 | SNOMED / ICD-10 ",
        foreign_key=(" DEMOGRAPHICS.PATIENT_HASH / LAB_TESTS.PATIENT_HASH "),
    )

    dictionary = import_dictionary(path)
    patient_hash = dictionary["DIAGNOSES"]["PATIENT_HASH"]

    assert patient_hash.vocabulary == ["ICD-10", "SNOMED", "ICD-10"]
    assert patient_hash.foreign_key == [
        "DEMOGRAPHICS.PATIENT_HASH",
        "LAB_TESTS.PATIENT_HASH",
    ]


@pytest.mark.parametrize(
    ("marker", "expected"), [("Y", True), ("y", True), ("", False)]
)
def test_import_parses_required_marker(tmp_path, marker: str, expected: bool):
    path = _dictionary_with_metadata_cells(
        tmp_path,
        vocabulary="ICD",
        foreign_key="DEMOGRAPHICS.PATIENT_HASH",
        required=marker,
    )

    dictionary = import_dictionary(path)

    assert dictionary["DIAGNOSES"]["PATIENT_HASH"].required is expected


@pytest.mark.parametrize("marker", ["N", "yes", "invalid"])
def test_import_rejects_non_y_required_markers(tmp_path, marker: str):
    path = _dictionary_with_metadata_cells(
        tmp_path,
        vocabulary="ICD",
        foreign_key="DEMOGRAPHICS.PATIENT_HASH",
        required=marker,
    )

    with pytest.raises(
        DataDictionaryImportError, match="Required must be 'Y' or blank"
    ):
        import_dictionary(path)


@pytest.mark.parametrize(
    ("foreign_key", "message"),
    [
        ("DEMOGRAPHICS.PATIENT_HASH / MISSING.PATIENT_HASH", "unknown table"),
        ("DEMOGRAPHICS.PATIENT_HASH / DEMOGRAPHICS.MISSING", "unknown column"),
        ("DEMOGRAPHICS.PATIENT_HASH / malformed", "must be 'TABLE.COLUMN'"),
    ],
)
def test_import_validates_each_foreign_key_target(tmp_path, foreign_key, message):
    path = _dictionary_with_metadata_cells(
        tmp_path, vocabulary="ICD", foreign_key=foreign_key
    )

    with pytest.raises(DataDictionaryImportError, match=message):
        import_dictionary(path)


if __name__ == "__main__":
    pytest.main([__file__])
